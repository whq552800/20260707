"""SDI等级×时期×疾病结局：稳健高尾与稳健低尾国家的因素构成（同期判定），以及同一等级同一时期内高尾与低尾的不利比例差。
因素与判定取自现行残差矩阵（MATRIX_FACTOR.csv：18项确认版因素，环境PM2.5为GBD来源，另加家庭PM2.5）。
不利比例 = 明确不利 /（明确不利+明确有利+方向不确定+已达普及）；有效年份不足与无数据不计入分母。任一组可判定国家<3时差值不计算。"""
import sys
from pathlib import Path
import numpy as np, pandas as pd, matplotlib
matplotlib.use('Agg'); import matplotlib.pyplot as plt
HERE = Path(__file__).resolve().parent; A = HERE.parent
sys.path[:0] = [str(A), str(A / 'figure2_reference_20261010' / 'qa_tools')]
from figure_export import apply_style, save_figure
from audit_panel_alignment import require_matplotlib_panel_alignment
from openpyxl.styles import PatternFill, Font, Alignment
apply_style({'font.family': ['Liberation Sans', 'Droid Sans Fallback'], 'font.size': 6.2, 'axes.linewidth': 0.4, 'xtick.major.size': 0, 'ytick.major.size': 0})
MM = 1 / 25.4; P3 = ['1990–2000', '2001–2010', '2011–2023']; TIERS = ['低SDI', '中低SDI', '中等SDI', '中高SDI', '高SDI']; MIN_N = 3
D = pd.read_csv(A / 'residual_matrix_workbook_20261010' / 'MATRIX_DISEASE.csv'); D = D[D.period.isin(P3)]
F = pd.read_csv(A / 'residual_matrix_workbook_20261010' / 'MATRIX_FACTOR.csv'); F = F[F.period.isin(P3)]
T = pd.read_csv(A / 'SDI_PERIOD_TIERS.csv'); loc = pd.read_csv(A / 'residual_matrix_workbook_20261010' / 'MATRIX_SDI.csv').drop_duplicates('location_id').set_index('location_id').location
D = D.merge(T[['location_id', 'period', 'tier']], on=['location_id', 'period']); D['outcome_cn'] = D.disease + D.outcome
D = D[D.category.isin(['稳健高尾', '稳健低尾'])]
FAC = list(dict.fromkeys(F.name)); OUTS = ['下呼吸道感染发病率', '下呼吸道感染死亡率', '脑膜炎发病率', '脑膜炎死亡率', '腹泻病发病率', '腹泻病死亡率']
JUDGE = ['明确不利', '明确有利', '方向不确定', '已达普及']
# 名单
names = D.groupby(['outcome_cn', 'period', 'tier', 'category']).location_id.apply(lambda s: '、'.join(sorted(loc[i] for i in s))).rename('countries').reset_index()
names['n'] = D.groupby(['outcome_cn', 'period', 'tier', 'category']).size().values
# 构成
m = D[['location_id', 'outcome_cn', 'period', 'tier', 'category']].merge(F[['location_id', 'period', 'name', 'status']], on=['location_id', 'period'], how='left')
grid = D[['location_id', 'outcome_cn', 'period', 'tier', 'category']].merge(pd.DataFrame({'name': FAC}), how='cross')
m = grid.merge(F[['location_id', 'period', 'name', 'status']], on=['location_id', 'period', 'name'], how='left'); m['status'] = m.status.fillna('无数据')
comp = m.groupby(['outcome_cn', 'period', 'tier', 'category', 'name']).status.value_counts().unstack(fill_value=0)
for c in JUDGE + ['有效年份不足', '无数据']:
    if c not in comp: comp[c] = 0
comp = comp[JUDGE + ['有效年份不足', '无数据']]; comp['可判定国家数'] = comp[JUDGE].sum(axis=1)
comp['不利比例%'] = (100 * comp['明确不利'] / comp['可判定国家数'].replace(0, np.nan)).round(1)
comp['有利比例%'] = (100 * comp['明确有利'] / comp['可判定国家数'].replace(0, np.nan)).round(1)
comp = comp.reset_index()
comp.to_csv(HERE / 'TIER_PERIOD_TAIL_COMPOSITION.csv', index=False, encoding='utf-8-sig')
# 差值
w = comp.pivot_table(index=['outcome_cn', 'period', 'tier', 'name'], columns='category', values=['不利比例%', '可判定国家数', '明确不利']).reset_index()
w.columns = ['outcome_cn', 'period', 'tier', 'name'] + [f'{a}|{b}' for a, b in w.columns[4:]]
for c in ['可判定国家数|稳健高尾', '可判定国家数|稳健低尾']:
    if c not in w: w[c] = 0
w = w.fillna({'可判定国家数|稳健高尾': 0, '可判定国家数|稳健低尾': 0})
ok = (w['可判定国家数|稳健高尾'] >= MIN_N) & (w['可判定国家数|稳健低尾'] >= MIN_N)
w['高尾减低尾（百分点）'] = np.where(ok, w['不利比例%|稳健高尾'] - w['不利比例%|稳健低尾'], np.nan)
w.to_csv(HERE / 'TIER_PERIOD_TAIL_DIFF.csv', index=False, encoding='utf-8-sig')
# 工作簿
out = HERE / 'SDI等级×时期×结局_稳健高低尾_因素构成_20261010.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as xw:
    names.rename(columns={'outcome_cn': '疾病结局', 'period': '时期', 'tier': 'SDI等级', 'category': '尾部', 'countries': '国家', 'n': '国家数'}).to_excel(xw, sheet_name='国家名单', index=False)
    comp.rename(columns={'outcome_cn': '疾病结局', 'period': '时期', 'tier': 'SDI等级', 'category': '尾部', 'name': '因素'}).to_excel(xw, sheet_name='因素构成', index=False)
    w.rename(columns={'outcome_cn': '疾病结局', 'period': '时期', 'tier': 'SDI等级', 'name': '因素'}).to_excel(xw, sheet_name='高尾与低尾差值', index=False)
    pd.DataFrame({'内容': ['分组', '因素判定', '不利比例', '差值', '解读'], '说明': [
        'SDI等级按GBD 2023固定参考区间、各时期SDI均值赋级；稳健高尾和稳健低尾为该时期、该疾病结局的现行分类（先区间后Q1/Q3）。',
        '使用同一时期的因素判定：明确不利、明确有利、方向不确定、已达普及（覆盖率类观测与参照均≥99%），以及有效年份不足（<2年）和无数据。环境PM2.5和家庭PM2.5均为GBD估计。',
        '明确不利国家数除以可判定国家数（明确不利、明确有利、方向不确定和已达普及之和）；有效年份不足和无数据不计入分母。',
        '同一结局、同一时期、同一SDI等级内，高尾组不利比例减低尾组不利比例；任一组可判定国家少于3个时不计算。正值表示该因素在高尾国家中更常处于不利。',
        '描述性汇总，各格国家数很少，未作检验；不表示因果或归因。']}).to_excel(xw, sheet_name='说明', index=False)
    for ws in xw.book.worksheets:
        for c in ws[1]: c.fill = PatternFill('solid', fgColor='FF1F4E78'); c.font = Font(bold=True, color='FFFFFFFF'); c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.freeze_panes = 'A2'
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 14
    xw.book['国家名单'].column_dimensions['F'].width = 90; xw.book['说明'].column_dimensions['B'].width = 110
# 图：每个结局一张热图
cmap = matplotlib.colormaps['RdBu_r']; norm = matplotlib.colors.TwoSlopeNorm(vmin=-100, vcenter=0, vmax=100)
SL = {'下呼吸道感染发病率': 'LRI_incidence', '下呼吸道感染死亡率': 'LRI_death', '脑膜炎发病率': 'meningitis_incidence', '脑膜炎死亡率': 'meningitis_death', '腹泻病发病率': 'diarrhoea_incidence', '腹泻病死亡率': 'diarrhoea_death'}
TS = {'低SDI': '低', '中低SDI': '中低', '中等SDI': '中等', '中高SDI': '中高', '高SDI': '高'}
for o in OUTS:
    x = w[w.outcome_cn == o]; cols = [(p, t) for p in P3 for t in TIERS]
    mean_diff = x.groupby('name')['高尾减低尾（百分点）'].mean()
    order = mean_diff.sort_values(ascending=False, na_position='last').index.tolist() + [f for f in FAC if f not in mean_diff.index]
    fig, ax = plt.subplots(figsize=(183 * MM, 132 * MM)); fig.subplots_adjust(left=0.16, right=0.90, top=0.90, bottom=0.08)
    for j, (p, t) in enumerate(cols):
        xj = j + (P3.index(p)) * 0.6
        for i, f in enumerate(order):
            r = x[(x.period == p) & (x.tier == t) & (x.name == f)]
            v = r['高尾减低尾（百分点）'].iloc[0] if len(r) else np.nan
            ax.add_patch(plt.Rectangle((xj - 0.46, len(order) - 1 - i - 0.46), 0.92, 0.92, facecolor='#F2F2F2' if np.isnan(v) else cmap(norm(v)), edgecolor='none'))
            if not np.isnan(v): ax.text(xj, len(order) - 1 - i, '0' if round(v) == 0 else f'{v:+.0f}', ha='center', va='center', fontsize=4.4, color='white' if abs(v) > 60 else '#222222')
        nh = len(D[(D.outcome_cn == o) & (D.period == p) & (D.tier == t) & (D.category == '稳健高尾')]); nl = len(D[(D.outcome_cn == o) & (D.period == p) & (D.tier == t) & (D.category == '稳健低尾')])
        ax.text(xj, -1.0, f'{nh}/{nl}', ha='center', va='center', fontsize=4.6, color='#444444')
        ax.text(xj, len(order) - 0.1, TS[t], ha='center', va='bottom', fontsize=5.4, rotation=0)
    for k, p in enumerate(P3):
        x0 = k * 5 + k * 0.6; ax.text(x0 + 2, len(order) + 0.9, p, ha='center', va='bottom', fontsize=6.6, fontweight='bold')
        ax.plot([x0 - 0.45, x0 + 4.45], [len(order) + 0.75] * 2, color='#444444', lw=0.5)
    ax.text(-0.9, -1.0, '高尾/低尾国家数', ha='right', va='center', fontsize=5.2, color='#444444')
    ax.set_xlim(-0.6, 15 + 1.2 - 0.4); ax.set_ylim(-1.6, len(order) + 1.8)
    ax.set_yticks(range(len(order))); ax.set_yticklabels([f.replace('₂.₅', '2.5') for f in order][::-1]); ax.set_xticks([])
    for sp in ax.spines.values(): sp.set_visible(False)
    cax = fig.add_axes([0.915, 0.30, 0.012, 0.40]); cb = fig.colorbar(matplotlib.cm.ScalarMappable(norm=norm, cmap=cmap), cax=cax)
    cb.set_label('高尾组减低尾组的不利比例（百分点）', fontsize=5.6); cb.ax.tick_params(labelsize=5.2); cb.outline.set_linewidth(0.3)
    fig.text(0.01, 0.985, f'{o}：同一SDI等级、同一时期内稳健高尾与稳健低尾国家的因素不利比例差', fontsize=7.4, fontweight='bold', ha='left', va='top')
    fig.text(0.01, 0.955, '红色表示该因素在高尾国家中更常处于明确不利；灰色为任一组可判定国家少于3个。因素按各格差值的均值排序。SDI等级按GBD 2023固定参考区间、各期SDI均值赋级；因素采用同期判定。',
             fontsize=5.4, ha='left', va='top', color='#444444', wrap=True)
    fig.canvas.draw(); save_figure(fig, HERE / f'FigS_tier_period_tail_{SL[o]}'); plt.close(fig)
print(out.name, 'rows', len(comp), len(w)); print(w.groupby('outcome_cn')['高尾减低尾（百分点）'].count().to_dict())
